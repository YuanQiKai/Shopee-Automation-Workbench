import json, math, sqlite3
from pathlib import Path
from datetime import date, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

base=Path(__file__).parent
rows=json.loads((base/'candidates.json').read_text(encoding='utf-8'))
wb=Workbook()
ws=wb.active; ws.title="使用说明"
notes=[
["项目","内容"],
["适用对象","菲律宾Shopee跨境店；不囤货；1688/义乌采购；供应商直发云仓；每天2–3小时"],
["资金与目标","启动20,000元；每月试错亏损上限4,000元；1–3个月验证品类，争取月净利>5,000元"],
["首轮结果","12个不同产品候选，0个可直接上架。所有候选均待补证，不代表建议下单。"],
["数据范围","2026-09-10查询7组PH商品关键词，140条不同商品；3组1688查询合计300条原始货源观测，未确认同款。"],
["价格缺口","市场价格原值禁止直接作PHP使用：PH查询与接口THB说明冲突。销量日期在搜索结果中为1970-01-01。"],
["货源缺口","3条供应商链接仅为线索；起价非对应套装报价；MOQ/代发/时效须核实。空白不是零。"],
["成本页","B:M全部输入人民币已核实值；确实不适用的项目填0并留证据。仅提示利润条件，不等于可发布。"],
["避免重复费用","商品净收入已扣卖家券；卖家券勿在成本再扣。物流扣款与补偿只入账一次。含税费勿重复加税。"],
["风险成本","成本页为预期测算；退款/拒收损失不能同时在净收入和风险成本重复扣。实际月结按真实归属核对。"],
["资金池","预算而非现金余额；应急5000元受保护。每月亏损上限不自动补充资金。"],
["流程","先核实费用、库存、包装、图权、交运，再审核；只有真实发布并回读可售才记上架完成。"],
["禁做","强资质认证、带电、化学相关、危险、违规、侵权、书籍、虚拟商品、收藏品。"],
["复盘","第7天查可售与曝光；14天分点击/转化；28天根据修正后的表现优化或停售，零曝光零单不直接删除。"],
["来源","市场及货源原始快照保存于同目录raw-observations.json；完整方案在artifact.json；自动执行规则在每日选品执行规则.md。"]
]
for r in notes: ws.append(r)
ws.column_dimensions['A'].width=20; ws.column_dimensions['B'].width=112

ws=wb.create_sheet("首轮12候选")
headers=["序号","产品方向","优先级","子方向","市场参考链接","差异化验证方向","关键检查","1688查找词","货源线索","货源说明","状态"]
ws.append(headers)
for r in rows:
    ws.append([r["rank"],r["name"],r["priority"],r["category"],r["market_url"],r["angle"],r["check"],r["sourcing"],r["source_supply"],r["supply_note"],r["status"]])
    for col in [5,9]:
        cell=ws.cell(ws.max_row,col)
        if cell.value: cell.hyperlink=cell.value;cell.style="Hyperlink"
widths=[8,30,16,16,28,44,52,32,28,40,26]
for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w

ws=wb.create_sheet("证据与待核实")
fields=[("rank","序号"),("name","产品"),("market_id","市场商品ID"),("price_raw","价格原值"),("price_unit","币种状态"),("monthly_raw","服务商月销量原值"),("monthly_note","销量状态"),("sales_calc_time","搜索销量日期"),("correction_date","服务商修正日期"),("shop_type","服务商店铺类型"),("shop_name","店铺名"),("brand_raw","品牌原字段待核实"),("supply_price_from","货源起价CNY非套装价"),("supply_moq_raw","MOQ原值"),("supply_drop_raw","代发原字段"),("observed_at","查询时间UTC")]
ws.append([label for _,label in fields])
for r in rows: ws.append([r[k] for k,_ in fields])
for i in range(1,len(fields)+1):ws.column_dimensions[get_column_letter(i)].width=24
ws.column_dimensions["B"].width=32;ws.column_dimensions["E"].width=48;ws.column_dimensions["G"].width=48

ws=wb.create_sheet("资金与情景")
for r in [
["项目","规划金额CNY","说明"],
["样品与验证",1500,"非销售库存"],
["首轮采购",5000,"出单采购"],
["国内物流/云仓/国际物流垫付",3000,"按实际付款时点计算"],
["测试推广",2000,"预算非授权自动花费"],
["工具运营",1000,"不自动充值"],
["滚动采购缓冲",2500,"与首轮采购合为7500"],
["保护应急金",5000,"不可用于新增采购"],
["总额","=SUM(B2:B8)","应为20000"],
["提前复盘亏损线",2000,"建议；按累计试错损失复盘"],
["每月亏损硬上限",4000,"用户确认；停止新增试错承诺"],
]:ws.append(r)
ws.append([])
ws.append(["净利CNY/单","月订单数（>5000）","按30天日均单量"])
for profit in [10,15,20,25]:
    n=5000//profit+1;ws.append([profit,n,n/30])
ws.append([])
ws.append(["回款天数","每日理论容量","假设说明"])
conn=sqlite3.connect(":memory:");conn.row_factory=sqlite3.Row
cash=[dict(r) for r in conn.execute((base/"cash-scenarios.sql").read_text(encoding="utf-8"))]
assert [r["capacity"] for r in cash]==[24,16,12]
for days,r in zip([14,21,28],cash):
    ws.append([days,r["capacity"],"采购池7500、物流池3000；单笔采购22、物流8元；未占用且无额外缓冲"])
ws.append(["注意","情景非预测","实际容量需扣已有承诺、退款等并核对每个资金池；月度亏损后重算。"])
for c,w in [("A",36),("B",26),("C",95)]:ws.column_dimensions[c].width=w

ws=wb.create_sheet("全成本核算待填写")
ws.append(["产品","商品净收入CNY","确认物流补偿CNY","采购","国内运费","云仓包材","卖家物流扣款","平台费用","广告推广","售后预计损失","收款汇兑","固定费用分摊","适用税费","全成本","预计净利","净利率","费用证据齐全","利润提示（不等于可发布）"])
for idx,r in enumerate(rows,2):
    ws.append([r["name"]]+[None]*15+["否",None])
    ws.cell(idx,14,f'=IF(COUNT(B{idx}:M{idx})<12,"",SUM(D{idx}:M{idx}))')
    ws.cell(idx,15,f'=IF(OR(COUNT(B{idx}:M{idx})<12,B{idx}<=0),"",B{idx}+C{idx}-N{idx})')
    ws.cell(idx,16,f'=IF(OR(COUNT(B{idx}:M{idx})<12,B{idx}<=0),"",O{idx}/B{idx})')
    ws.cell(idx,18,f'=IF(OR(Q{idx}<>"是",COUNT(B{idx}:M{idx})<12),"待补证",IF(B{idx}<=0,"收入无效",IF(AND(O{idx}>0,P{idx}>=0.2),"利润条件达标；仍需其他审核","利润不达标")))')
    for col in range(2,16):ws.cell(idx,col).number_format='0.00'
    ws.cell(idx,16).number_format="0.0%"
dv=DataValidation(type="list",formula1='"是,否"',allow_blank=False);ws.add_data_validation(dv);dv.add("Q2:Q13")
for i in range(1,19):ws.column_dimensions[get_column_letter(i)].width=20
ws.column_dimensions["A"].width=34;ws.column_dimensions["R"].width=45
ws.append(["说明","仅为单笔预期测算模板，非平台计费引擎。各费用需按真实基数与封顶算出后填入；费用缺项时冻结结论。"])

ws=wb.create_sheet("30天日报")
ws.append(["日期","新增候选目标","实际新增候选","证据完整","真实上架成功","复查数量","当日归属净利CNY","实际可用现金CNY","待补证/失败/次日动作"])
for day in range(30):ws.append([date(2026,9,10)+timedelta(days=day),10,None,None,None,None,None,None,None])
for cell in ws["A"][1:]:cell.number_format="yyyy-mm-dd"
for i in range(1,10):ws.column_dimensions[get_column_letter(i)].width=24
ws.column_dimensions["I"].width=80

for ws in wb:
    ws.freeze_panes="B2";ws.auto_filter.ref=ws.dimensions
    ws.sheet_view.showGridLines=False
    for cell in ws[1]:
        cell.font=Font(name="微软雅黑",bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="163A4A")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment=Alignment(vertical="top",wrap_text=True)
            if not cell.hyperlink:cell.font=Font(name="微软雅黑",size=11)
            if cell.row%2==0:cell.fill=PatternFill("solid",fgColor="F1F5F7")
    ws.row_dimensions[1].height=32
    for i in range(2,ws.max_row+1):ws.row_dimensions[i].height=52
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation="landscape";ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
out=base/"菲律宾Shopee每日选品执行表.xlsx"
wb.save(out)
check=load_workbook(out,data_only=False)
assert check["首轮12候选"].max_row==13
assert len({check["首轮12候选"].cell(r,5).value for r in range(2,14)})==12
assert sum([1500,5000,3000,2000,1000,2500,5000])==20000
assert all(check["全成本核算待填写"].cell(r,2).value is None for r in range(2,14))
assert all(check["全成本核算待填写"].cell(r,18).data_type=="f" for r in range(2,14))
print(json.dumps({"file":str(out),"sheets":check.sheetnames,"candidates":12,"verified_publishable":0,"cash_capacity":[r["capacity"] for r in cash],"qa":"passed; formulas written, no spreadsheet-engine recalculation"},ensure_ascii=False))

