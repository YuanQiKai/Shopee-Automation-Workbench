"""Isolated PostgreSQL schema; no user candidates or credentials are changed."""
import sys,uuid,io,copy,json
from pathlib import Path
from datetime import datetime,timezone,timedelta
from decimal import Decimal as D
import pytest
from sqlalchemy import create_engine,text
from sqlalchemy.orm import sessionmaker
from fastapi.testclient import TestClient
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import backend.app as api
from backend.db import Base,db_url,Record,put
from backend.domain import evaluate,current_hash,supplier_view,GATES
from test_finance import scenario
@pytest.fixture
def client(monkeypatch):
    schema='qa_'+uuid.uuid4().hex
    admin=create_engine(db_url(),connect_args={'connect_timeout':5})
    with admin.begin() as c:c.execute(text('CREATE SCHEMA '+schema))
    engine=create_engine(db_url(),connect_args={'options':'-csearch_path='+schema,'connect_timeout':5})
    sessions=sessionmaker(engine,expire_on_commit=False)
    monkeypatch.setattr(api,'engine',engine);monkeypatch.setattr(api,'Session',sessions)
    with TestClient(api.app,headers={'Origin':'http://127.0.0.1:3010','X-Meeya-Local':'1'}) as c:yield c,sessions
    engine.dispose()
    assert schema.startswith('qa_') and len(schema)==35
    with admin.begin() as c:c.execute(text('DROP SCHEMA '+schema+' CASCADE'))
    admin.dispose()
def test_seed_export_and_unknown_blank(client):
    c,s=client;d=c.get('/api/bootstrap').json();assert len(d['candidates'])==22
    assert sum(len(x['skus']) for x in d['suppliers'])==38
    assert all(not x['evaluation']['eligible'] for x in d['candidates'])
    r=c.get('/api/export');assert r.status_code==200
    wb=load_workbook(io.BytesIO(r.content),data_only=True);assert len(wb.sheetnames)==10
    assert wb['01审核清单'].max_row==23 and wb['01审核清单']['L2'].value is None
def test_optimistic_lock_and_reload(client):
    c,s=client;d=c.get('/api/bootstrap').json();first=d['candidates'][0]
    body={'revision':first['revision'],'data':{'spec':'隔离测试修改'}}
    assert c.put('/api/candidates/C01',json=body).status_code==200
    assert c.put('/api/candidates/C01',json=body).status_code==409
    with s() as session:assert session.get(Record,('candidate','C01')).data['spec']=='隔离测试修改'
def test_cross_origin_and_approval_block(client):
    c,s=client
    assert c.post('/api/candidates',json={'name':'bad'},headers={'Origin':'https://example.org'}).status_code==403
    d=c.get('/api/bootstrap').json()['candidates'][0]
    assert c.post('/api/candidates/C01/review',json={'state':'approved','revision':d['revision'],'note':'隔离测试'}).status_code==422
def test_secret_never_returned_and_encrypted(client):
    c,s=client;secret='qa-only-not-a-real-key'
    r=c.post('/api/providers',json={'name':'隔离测试','protocol':'sorftime','api_key':secret})
    assert r.status_code==200 and secret not in r.text and 'secret' not in r.json()
    assert secret not in c.get('/api/bootstrap').text and secret not in c.get('/api/audit').text
    with s() as session:
        stored=session.get(Record,('provider',r.json()['id'])).data['secret'];assert stored!=secret
def fixture_ready():
    c,settings=scenario();now=datetime.now(timezone.utc);at=now.isoformat();end=(now.date()-timedelta(days=1));start=end-timedelta(days=29);expiry=(now+timedelta(days=1)).isoformat()
    settings.update(fee_evidence='隔离测试',fee_valid_until=expiry,warehouse_evidence='隔离测试',available_cash_cny='5100',commitments_cny='0',reserve_cny='5000',monthly_loss_limit_cny='4000',realized_loss_cny='0',unresolved_loss_cny='0',cash_as_of=now.astimezone().date().isoformat())
    c.update(id='QA1',name='隔离测试',planned_landed_php='440',difference_evidence='隔离测试',market_observations=[dict(product_id=str(i),shop_id=str(i//2),source='隔离测试',site='PH',currency='PHP',spec=c['spec'],sku_sales_verified=True,sales_kind='服务商估算',sales_30d=100,buyer_landed_php=440,shop_type='跨境' if i<4 else '本土',window_start=str(start),window_end=str(end),observed_at=at) for i in range(10)],checks={k:{'status':'pass','evidence':'隔离测试','valid_until':expiry} for k in GATES})
    c['financial']['buyer_shipping_php']='40'
    suppliers=[]
    for i,role in enumerate(('primary','backup')):
        id=str(i);suppliers.append({'offer_id':id,'captured_at':at,'detail_at':at,'skus_at':at,'detail':{'store_name':'测试商家'+id},'skus':[{'sku_id':'s'+id,'stock':100,'price':5,'offer_price':5,'length':12,'width':8,'height':1,'weight':.1}]})
        c[role]={'offer_id':id,'sku_id':'s'+id,'verification':{'same_spec':True,'quote_evidence':'隔离测试','pack_measured':True,'confirmed_at':at,'lead_hours':24,'moq':1,'dropship':True,'final_price_cny':'5'}}
    return c,settings,suppliers
def test_hash_matches_annotated_export():
    c,s,sups=fixture_ready();a=current_hash(c,s,sups)
    b=current_hash(c,s,[supplier_view(copy.deepcopy(x)) for x in sups]);assert a==b
    c['financial']['purchase_price_cny']='6';assert current_hash(c,s,sups)!=a
def test_joint_cash_prevents_double_spend_and_stale(client):
    c,sessions=client;candidate,settings,suppliers=fixture_ready();e=evaluate(candidate,settings,suppliers);assert e['eligible'],e['reasons']
    with sessions.begin() as s:
        put(s,'settings','shop',settings)
        for sup in suppliers:put(s,'supplier',sup['offer_id'],sup)
        for i in ('QA1','QA2'):put(s,'candidate',i,{**copy.deepcopy(candidate),'id':i})
    assert c.post('/api/candidates/QA1/review',json={'revision':1,'state':'approved','note':'隔离测试'}).status_code==200
    r=c.post('/api/candidates/QA2/review',json={'revision':1,'state':'approved','note':'隔离测试'});assert r.status_code==422 and '现金' in r.text
    d=next(x for x in c.get('/api/bootstrap').json()['candidates'] if x['id']=='QA1')
    assert d['evaluation']['review_state']=='approved'
    assert c.put('/api/candidates/QA1',json={'revision':d['revision'],'data':{'spec':'已变化'}}).status_code==200
    d=next(x for x in c.get('/api/bootstrap').json()['candidates'] if x['id']=='QA1');assert d['evaluation']['review_state']=='stale'
