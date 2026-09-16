import io,json
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from .costs import RATES
def safe(v):
    if isinstance(v,(dict,list)):v=json.dumps(v,ensure_ascii=False,default=str)
    if isinstance(v,str) and v[:1] in '=+-@':v="'"+v
    return v
def workbook(candidates,settings,suppliers,evidence,evaluations):
    wb=Workbook();wb.remove(wb.active)
    def sheet(name,headers,rows):
        ws=wb.create_sheet(name);ws.append(headers)
        for row in rows:ws.append([safe(v) for v in row])
        ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.fill=PatternFill('solid',fgColor='204D43');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(65,max(15,len(str(col[0].value))*2+5))
            for cell in col[1:]:cell.alignment=Alignment(vertical='top',wrap_text=True)
        for i in range(2,ws.max_row+1):ws.row_dimensions[i].height=45
        return ws
    summary=[];cost=[];gates=[]
    num=lambda x:float(Decimal(x)) if x is not None else None
    for c in candidates:
        v=evaluations[c['id']];b=v['base'];st=v['stress']
        summary.append([c['id'],c['name'],v['review_state'],'可申请审核' if v['eligible'] else '待补证/未通过',c.get('spec'),c.get('fact'),c.get('inference'),c.get('recommendation'),num(b['sale_php']),num(b['revenue_cny']),num(b['full_cost_cny']),num(b['profit_cny']),num(b['margin_pct']),num(b['purchase_ceiling_cny']),num(st['profit_cny']),num(b['test_total_cny']),'; '.join(v['reasons']),c.get('market_url'),c.get('primary'),c.get('backup'),v['input_hash']])
        for scenario in ('base','stress'):
            for line in v[scenario]['lines']:cost.append([c['id'],v[scenario]['scenario'],line['name'],num(line['cny']),num(line.get('php')),'缺值' if line['cny'] is None else '计算/预算；须核实来源'])
        for row in c.get('evidence_rows',[]):gates.append([c['id'],*row])
        for k,x in c.get('checks',{}).items():gates.append([c['id'],k,x.get('status'),x.get('valid_until'),x.get('evidence'),'人工补证',''])
    sheet('01审核清单',['ID','候选产品','审核状态','可做程度','规格','事实','推断','建议','活动售价PHP','收入CNY','完全成本CNY','利润CNY','利润率%','采购单价上限CNY','压力利润CNY','测试累计总成本CNY','阻塞原因','市场链接','主货源','备用货源','输入版本'],summary)
    sheet('02成本逐项',['产品ID','场景','费项','金额CNY','原金额PHP','证据状态'],cost)
    sheet('03证据与门槛',['产品ID','维度','实际值/状态','阈值/期限','符合程度/来源','原因','证据位置'],gates)
    sheet('04供货SKU',['offer ID','商家','SKU ID','规格','展示价CNY','报价字段CNY','库存','长cm','宽cm','高cm','重kg','包装来源','采集时间','异常','证据ID'],[[sup['offer_id'],sup.get('detail',{}).get('store_name'),str(k['sku_id']),k.get('sku_name'),k.get('price'),k.get('offer_price'),k.get('stock'),k.get('length'),k.get('width'),k.get('height'),k.get('weight'),k.get('pkg_size_source'),sup.get('captured_at') or sup.get('captured_date'),k.get('issues'),sup.get('evidence_ids')] for sup in suppliers for k in sup.get('skus',[])])
    sheet('05店铺资金费率',['字段','值','说明'],[[k,v,'空白保持未知；税费不可用单一推测率'] for k,v in settings.items()])
    sheet('06政策来源',['编号','来源','位置','定位','资料日期','核查日期','结论'],[[r['id'],r['title'],r['location'],r.get('locator'),r.get('date'),r.get('checked'),r['finding']] for r in RATES['sources']])
    sheet('07海运禁运',['类目ID','类目路径','原行','无标题备注','判定说明'],[[b['category_id'],b['path'],b['row'],b['raw_note'],b['note']] for b in RATES['sea_ban']])
    sheet('08采集记录',['证据ID','接口','参数','采集时间','原始结果位置'],[[x['id'],x['tool'],x['arguments'],x.get('captured_at') or x.get('captured_date'),'/api/evidence/'+x['id']] for x in evidence])
    sheet('09计算输入',['产品ID','字段','输入值'],[[c['id'],k,v] for c in candidates for k,v in c.get('financial',{}).items()])
    sheet('10使用说明',['事项','说明'],[['利润定义','一销售单元/一件平台商品/一订单；不同件数组合须重新核实费用；完整成本及压力计算在Python Decimal执行，导出保留输入和输出快照。'],['未知数据','空白不等于0；未经确认的来源/预测不得称为盈利事实。'],['活动','4折=标价×0.4的演算，未确认承担方前不能通过。'],['供应商','接口报价与展示价不一致、包装0、城市型时效均保留；主备需同规格及最终报价确认。'],['本次验收','没有足够证据时如实报告0个可做；不凑满10个。']])
    out=io.BytesIO();wb.save(out);return out.getvalue()
