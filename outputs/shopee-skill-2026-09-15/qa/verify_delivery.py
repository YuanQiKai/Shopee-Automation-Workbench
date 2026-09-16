from pathlib import Path
import sys,json,hashlib,contextlib,io,copy
import openpyxl
BASE=Path(__file__).resolve().parent.parent
ROOT=BASE.parent.parent
sys.path.insert(0,str(BASE/'shopee-ph-selection/scripts'))
sys.path.insert(0,str(ROOT/'outputs/selection-sop-review-2026-09-14/qa/calcdeps'))
import formulas
from build_daily import build
from test_selection_math import fixture,NOW,run
from selection_math import calculate,GATES

checks=[];checks.append({'check':'Decimal利润/压力/缺失值/边界','result':run()})
data=json.loads((BASE/'selection.json').read_text(encoding='utf-8'))
for x in json.loads((BASE/'evidence/input-manifest.json').read_text(encoding='utf-8')):
    assert hashlib.sha256((ROOT/x['source']).read_bytes()).hexdigest()==x['sha256']
    assert hashlib.sha256((BASE/x['snapshot']).read_bytes()).hexdigest()==x['sha256']
checks.append({'check':'原始源文件与快照SHA256','result':'一致，未改源文件'})
actual=BASE/'菲律宾Shopee每日选品审核.xlsx'
b=openpyxl.load_workbook(actual,data_only=True)
assert len(b.sheetnames)==13
assert b['01_逐品结论'].max_row==26
assert b['02_逐项证据门槛'].max_row==4+22*9
assert b['04_费用明细'].max_row==4+22*18
assert b['06_1688历史货源'].max_row==604
assert b['07_历史市场原值'].max_row==224
for r in range(5,27):
    assert b['01_逐品结论'].cell(r,3).value=='待补证'
    for col in [11,12,13]:assert b['01_逐品结论'].cell(r,col).value is None
    assert b['03_逐品成本'].cell(r,20).value in [None,'']
    assert b['03_逐品成本'].cell(r,26).value=='待补证'
checks.append({'check':'真实交付覆盖/缺值安全','result':'13表、22候选、198维度核验、396费用行；利润未填0，全部待补证'})

def calc_book(model,name):
    d=copy.deepcopy(data);d['as_of']=NOW;d.pop('flow_png',None)
    c=copy.deepcopy(data['candidates'][0]);c.update(id='TEST',name='隔离合成算例，非市场产品',model=model,gates={k:'通过' for k in GATES})
    d['candidates']=[c];d['observations']=[];d['suppliers']=[];d['issues']=[]
    p=BASE/'qa'/f'{name}.xlsx';build(d,p)
    book=openpyxl.load_workbook(p)
    for sh in list(book.sheetnames):
        if sh not in ['00_审核入口','03_逐品成本','04_费用明细']:del book[sh]
    book.save(p)
    with contextlib.redirect_stdout(io.StringIO()),contextlib.redirect_stderr(io.StringIO()):
        result=formulas.ExcelModel().loads(str(p)).finish().calculate()
    def value(sheet,cell):
        k=next(k for k in result if k.upper().endswith(sheet.upper()+"'!"+cell.upper()))
        return result[k].value[0,0]
    return value

m=fixture();v=calc_book(m,'numeric-formulas');expected=calculate(m,NOW)
for cell,key in [('O5','other_cny'),('P5','stress_other_cny'),('Q5','revenue_cny'),('R5','procurement_cny'),('S5','total_cost_cny'),('T5','profit_cny'),('U5','margin'),('V5','quote_ceiling_cny'),('W5','stress_profit_cny'),('X5','test_total_cny'),('Y5','max_test_orders')]:
    assert abs(float(v('03_逐品成本',cell))-float(expected[key]))<1e-8,(cell,v('03_逐品成本',cell),expected[key])
assert v('03_逐品成本','Z5')=='成本条件通过'
checks.append({'check':'实际Excel公式跨引擎重算','result':'11个成本/压力/反向报价/预算单元格与Decimal一致'})
m=fixture();m['setup_cny']=150;v=calc_book(m,'budget-exact-500');assert float(v('03_逐品成本','X5'))==500 and v('03_逐品成本','Z5')=='成本条件通过'
m=fixture();m['setup_cny']=151;v=calc_book(m,'budget-over-500');assert v('03_逐品成本','Z5')=='成本条件不通过'
m=fixture();m['fee_lines'][0]['amount']=None;v=calc_book(m,'missing-cost');assert v('03_逐品成本','Z5')=='待补证';assert v('03_逐品成本','T5')==''
checks.append({'check':'Excel实际边界行为','result':'500元可通过、501元不通过、缺费项保持待补证且利润空白'})
payload={'checks':checks,'limits':['未在本机Excel/WPS图形界面复算；已用独立Excel公式引擎实际计算隔离样本。','真实供应链/结算对账与盈利测试尚未执行。','流程图已渲染并目视检查；Notebook代码通过Python执行，未在Jupyter内核执行。']}
(BASE/'qa/verification-results.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(payload,ensure_ascii=False,indent=2))
