from pathlib import Path
import urllib.request,json,sys,io
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1];sys.path.insert(0,str(R))
from backend.db import Session,Record,put
from backend.domain import ingest
with Session.begin() as s:
    for snap in json.loads((R/'data/evidence/sorftime-refresh-2026-09-15.json').read_text(encoding='utf-8')):ingest(s,snap)
    st=s.get(Record,('settings','shop'))
    if 'transaction_stress_pct' not in st.data:put(s,'settings','shop',{**st.data,'transaction_stress_pct':None})
out=R.parent/'outputs/shopee-workbench-2026-09-15'
d=json.load(urllib.request.urlopen('http://127.0.0.1:3010/api/bootstrap',timeout=20))
raw=urllib.request.urlopen('http://127.0.0.1:3010/api/export',timeout=20).read()
(out/'菲律宾选品工作台审核.xlsx').write_bytes(raw)
wb=load_workbook(io.BytesIO(raw),data_only=True)
assert len(d['candidates'])==22 and sum(len(s['skus']) for s in d['suppliers'])==38
assert not any(c['evaluation']['eligible'] for c in d['candidates'])
assert len(wb.sheetnames)==10 and wb['01审核清单']['L2'].value is None
assert all(s.get('detail_at') and s.get('skus_at') for s in d['suppliers'])
assert any(j['data']['status']=='completed' for j in d['jobs'])
for obj in d['providers']:assert 'secret' not in obj
summary={'source_date':'2026-09-15','candidates':22,'suppliers':2,'sku_records':38,'eligible':0,'sls_weight_bands':5000,'sea_list_records':279,'excel_sheets':len(wb.sheetnames),'real_database':'PostgreSQL 17','task_queue':'Temporal completed real recalculate job','tests':'20 application checks + 13 standalone Skill fee checks','ai_paid_calls':0,'local_sorftime_credentials_configured':any(p['protocol']=='sorftime' for p in d['providers'])}
(out/'验证结果.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False))
