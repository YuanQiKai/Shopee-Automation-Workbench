from pathlib import Path
from pptx import Presentation
from docx import Document
import json,hashlib
root=Path(__file__).resolve().parents[1]
paths=[Path('G:/Shopee/货代/源宇云仓简介和收费标准等.pptx'),Path('G:/Shopee/货代/重要必看合作配合流程.docx')]
out=[]
for p in paths:
    texts=[]
    if p.suffix=='.pptx':
        for i,s in enumerate(Presentation(p).slides):
            lines=[]
            for sh in s.shapes:
                if sh.has_text_frame: lines.append(sh.text)
                if sh.has_table: lines += [' | '.join(c.text for c in r.cells) for r in sh.table.rows]
            texts.append({'slide':i+1,'text':'\n'.join(lines)})
    else:
        d=Document(p)
        texts=[{'text':t.text} for t in d.paragraphs if t.text]
        texts += [{'table':i+1,'rows':[[c.text for c in row.cells] for row in t.rows]} for i,t in enumerate(d.tables)]
    out.append({'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'content':texts})
f=root/'data/evidence/warehouse-documents.json'
f.write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(out,ensure_ascii=False,indent=2))
