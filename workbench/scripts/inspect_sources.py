from pathlib import Path
import json, hashlib
from pypdf import PdfReader
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
out = root / 'data' / 'evidence'
out.mkdir(parents=True, exist_ok=True)
files = [Path('G:/Shopee/物流手册.pdf'), Path('G:/Shopee/菲律宾海运禁运.xlsx'), Path('G:/Shopee/常用表格/跨境物流成本（藏价）计算工具 - V2.0 20260908.xlsx')]
manifest=[]
for p in files:
    record={'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'size':p.stat().st_size}
    if p.suffix=='.pdf':
        pages=[{'page':i+1,'text':page.extract_text()} for i,page in enumerate(PdfReader(p).pages)]
        (out/'logistics-manual-pages.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2),encoding='utf-8')
        record['pages']=len(pages)
        print(json.dumps({'pdf':str(p),'pages':len(pages),'ph_pages':[x['page'] for x in pages if '菲律宾' in x['text']]},ensure_ascii=False))
    else:
        wb=load_workbook(p,read_only=True,data_only=True)
        record['sheets']=[{'name':s.title,'rows':s.max_row,'cols':s.max_column} for s in wb]
        result={}
        for s in wb:
            if '禁运' in p.name or any(t in s.title for t in ['菲律宾','SLS','运费','说明']):
                rows=[]
                for cells in s:
                    values={c.column_letter:c.value for c in cells if c.value is not None}
                    if values: rows.append({'row':cells[0].row,'cells':values})
                result[s.title]=rows
        (out/('sea-ban-rows.json' if '禁运' in p.name else 'sls-workbook-rows.json')).write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
        print(json.dumps(record,ensure_ascii=False))
        wb.close()
    manifest.append(record)
(out/'input-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
